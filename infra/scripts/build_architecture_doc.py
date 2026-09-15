"""
The System Architecture Document, built from the code and the pictures.

    npm run docs:architecture        (after npm run docs:diagrams and docs:shots)

One Word document, and a PDF beside it when Word is present: how the system
is put together, how a request is processed, how data is written and moved,
the database domain by domain with an entity-relationship diagram read from
the Prisma schema, the security in the data path, the user interface as
wireframes and photographs of the running app, and what verifies it.

Rebuilt after every change, with the design book and the deck. Nothing in
it is typed twice: counts and tables come from docs/assets/architecture/
schema.json, diagrams from the schema, screenshots from design-shots.ts.
"""

from __future__ import annotations

import json
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
ASSETS = ROOT / "docs" / "assets" / "architecture"
SHOTS = ROOT / "docs" / "assets" / "design"
WIREFRAMES = ROOT / "docs" / "02-design" / "09-wireframes.md"
MIGRATIONS = ROOT / "packages" / "db" / "prisma" / "migrations"
SCRIPTS = ROOT / "infra" / "scripts"
OUT = ROOT / "docs" / "Carnet-System-Architecture.docx"

DISPLAY = "Archivo"
BODY = "Public Sans"
MONO = "IBM Plex Mono"
INK = RGBColor(0x1A, 0x20, 0x22)
PETROL = RGBColor(0x1F, 0x61, 0x67)
PETROL_DEEP = RGBColor(0x12, 0x32, 0x36)
BRASS = RGBColor(0x9A, 0x6D, 0x24)
FAINT = RGBColor(0x6B, 0x75, 0x77)

BUILT_ON = "14 September 2026"
AFTER_PHASE = "Phase 35"

PORTRAIT_W = 16.6
LANDSCAPE_W = 25.3


# ----------------------------------------------------------------- helpers
def style_document(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = BODY
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.18
    for name, size, colour, before in (
        ("Heading 1", 22, PETROL_DEEP, 18),
        ("Heading 2", 15, PETROL, 14),
        ("Heading 3", 12, INK, 10),
    ):
        style = doc.styles[name]
        style.font.name = DISPLAY
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = colour
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    portrait(doc.sections[0])


def portrait(section) -> None:
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.2)
    section.top_margin = section.bottom_margin = Cm(2.0)


def landscape(section) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.left_margin = section.right_margin = Cm(2.2)
    section.top_margin = section.bottom_margin = Cm(1.8)


def para(doc, text: str, *, size=10.5, colour=INK, bold=False, italic=False, font=BODY,
         align=None, space_after=8, keep=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = colour
    run.bold = bold
    run.italic = italic
    p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    if keep:
        p.paragraph_format.keep_with_next = True
    return p


def label(doc, text: str):
    return para(doc, text.upper(), size=8.5, colour=BRASS, bold=True, space_after=2, keep=True)


def bullets(doc, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        run.font.name = BODY
        run.font.size = Pt(10.5)
        p.paragraph_format.space_after = Pt(3)


def mono(doc, text: str, size=8.0):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = MONO
    run.font.size = Pt(size)
    run.font.color.rgb = INK
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.0
    pPr = p._p.get_or_add_pPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:val"), "clear")
    shade.set(qn("w:fill"), "F4F1EA")
    pPr.append(shade)
    return p


def image(doc, path: Path, caption: str, width_cm: float):
    if not path.exists():
        para(doc, f"[missing picture: {path.name}]", italic=True, colour=FAINT)
        return
    doc.add_picture(str(path), width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, caption, size=9, colour=FAINT, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)


def fit(path: Path, max_w: float, max_h: float) -> float:
    """Width in cm that keeps the picture inside max_w x max_h."""
    from PIL import Image  # pillow ships with python-docx's environment here

    with Image.open(path) as im:
        w, h = im.size
    width = max_w
    if h / w * width > max_h:
        width = max_h * w / h
    return width


def diagram(doc, name: str, caption: str, max_w: float = PORTRAIT_W, max_h: float = 20.0):
    path = ASSETS / f"{name}.png"
    if not path.exists():
        para(doc, f"[missing diagram: {name}.png — run architecture-diagrams.ts]", italic=True, colour=FAINT)
        return
    image(doc, path, caption, fit(path, max_w, max_h))


def shot(doc, name: str, caption: str, width_cm: float = PORTRAIT_W):
    image(doc, SHOTS / f"{name}.png", caption, width_cm)


def table(doc, header: list[str], rows: list[list[str]], widths: list[float] | None = None, size=8.5):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(size)
        run.font.name = BODY
        run.font.color.rgb = PETROL_DEEP
        shade = OxmlElement("w:shd")
        shade.set(qn("w:val"), "clear")
        shade.set(qn("w:fill"), "E6F0F1")
        cell._tc.get_or_add_tcPr().append(shade)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(value)
            run.font.size = Pt(size)
            run.font.name = MONO if i == 0 else BODY
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def rule(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "D9B36C")
    border.append(bottom)
    pPr.append(border)
    p.paragraph_format.space_after = Pt(4)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def workflows_table(doc) -> None:
    """The press-count tables from docs/02-design/11-workflows.md, which is generated from core."""
    path = ROOT / "docs" / "02-design" / "11-workflows.md"
    if not path.exists():
        return
    title = None
    rows: list[list[str]] = []
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.startswith("### "):
            if title and rows:
                label(doc, title)
                table(doc, ["Hand-off", "Screen", "Before", "Now"], rows, widths=[6.0, 2.4, 1.4, 1.4], size=8.0)
            title = line[4:].strip()
            rows = []
        elif line.startswith("| ") and not line.startswith("| ---") and not line.startswith("| Hand-off"):
            cells = [c.strip().replace("**", "") for c in line.strip().strip("|").split("|")]
            if len(cells) >= 4:
                rows.append(cells[:4])
    if title and rows:
        label(doc, title)
        table(doc, ["Hand-off", "Screen", "Before", "Now"], rows, widths=[6.0, 2.4, 1.4, 1.4], size=8.0)


def wireframe(section_title: str) -> str | None:
    """The first code block under a heading in the wireframes document."""
    text = WIREFRAMES.read_text(encoding="utf-8")
    m = re.search(r"^## [^\n]*" + re.escape(section_title) + r"[^\n]*\n(.*?)(?=^## |\Z)", text, re.M | re.S)
    if not m:
        return None
    block = re.search(r"```\n(.*?)```", m.group(1), re.S)
    return block.group(1).rstrip() if block else None


# ----------------------------------------------------------------- content
DOMAIN_TEXT: dict[str, list[str]] = {
    "tenancy": [
        "Everything hangs off the organisation. A workshop is an Organization with one or more Locations, each with Bays; "
        "a person is a User with a Membership that carries their role in one organisation. Sessions are rows, not signed "
        "tokens: the cookie holds a random token whose SHA-256 digest is the row's key, so a session can be ended on the "
        "server and a stolen database does not yield usable cookies (ADR-0006).",
        "The organisation also carries its region (country, currency, locale, timezone) and the four dates its subscription "
        "state is read from. AuditLog is written by every consequential action and never rolls the action back if it fails; "
        "OutboxEvent is the transactional outbox that carries side effects out of the commit that caused them; SyncedAction "
        "is the idempotency record for what a technician's device did while it was out of signal. ImportBatch is a file "
        "brought from another system, kept as it arrived with its mapping and its report.",
    ],
    "vehicle": [
        "The vehicle is global, not owned by a workshop (ADR-0004): one Vehicle row per VIN, whoever touches it. Its "
        "history is VehicleEvent, append-only and hash-chained: each event's hash covers its own content and the hash of "
        "its predecessor, so altering any event breaks every hash after it, and a correction is a new event that supersedes "
        "an old one rather than an edit. VehicleOwnership is the chain of custody; OdometerReading is kept separately so "
        "a reading can be checked against every other.",
        "Who may see what is VehicleAccessGrant (a workshop's entitlement, earned by working on the car or granted by "
        "the owner) and PassportShare (a link with a scope and an expiry, every opening recorded in PassportShareAccess). "
        "VehicleDocument is the paperwork; VehicleClaim is a warranty or insurance claim against the record; Recall and "
        "VehicleRecall are manufacturer notices matched to vehicles; PartFitment is what fits what.",
    ],
    "customer": [
        "A Customer belongs to one organisation and may be an individual or a company; a fleet is a company Customer with "
        "a FleetAgreement that sets the terms for every car on it. Booking is a diary slot against a bay and a capacity "
        "check; JobType and ServiceItem are the workshop's own catalogue of what it does and what it charges; "
        "ServiceSchedule is the next service the reminder job will write to the customer about.",
        "ConsentChange is the record of every change to what a customer has agreed to be sent, which is what lets the "
        "product answer a complaint with a date rather than an apology.",
    ],
    "job": [
        "The Job is the central record of the workshop. It carries its own lines (labour, parts, fees) rather than "
        "pointing at an estimate's, because what was quoted and what was done are allowed to differ and both are kept. "
        "Every status change is a JobStatusEvent, so the job's timeline is a query, not a memory. TimeEntry is the "
        "technician's clock; QualityCheck is the sign-off before the customer is told; JobFollowUp is the call afterwards.",
        "Totals are never stored on the job: they are derived from the lines each time they are read (ADR-0011), which is "
        "why a line changing cannot leave a stale figure behind.",
    ],
    "inspection": [
        "An InspectionTemplate is the workshop's own checklist, in sections and items; an Inspection is one run of it "
        "against one job, each item answered in an InspectionResult with a rating and, where it matters, an "
        "InspectionPhoto. The photograph is the trust mechanism (ADR-0012): a customer approves brake work because they "
        "can see the disc, not because they are told about it.",
        "An Estimate is built from the results, line by line, with the currency of the organisation on it. It goes to "
        "the customer as an EstimateShare, a link with a token, and comes back as an EstimateApproval that records who "
        "approved which lines, when, from which address. An approved estimate is what a job's work is authorised against.",
    ],
    "invoice": [
        "An Invoice is raised from a job with its lines copied and its tax computed by the TaxRate in force; its number "
        "comes from DocumentSequence, which hands out gapless numbers per organisation and per document kind under a "
        "row lock. Totals are Decimal(12,2) in the database and integer minor units in the rules, converted at the "
        "boundary, so arithmetic never happens in floating point.",
        "A Payment is one tender against one invoice, with a method that decides its ledger account: cash and card to "
        "the bank, points to the loyalty liability, a gift card to the gift card liability, money held to the deposit "
        "liability. A CreditNote reverses an invoice in whole or part; a Refund is money going back out, with a method "
        "of its own, and may instead be kept as credit. A PaymentAttempt is one try at paying through the provider, "
        "moving once from pending; a succeeded attempt produces exactly one Payment, and a Refund sent back through "
        "the provider names the attempt it went back to.",
    ],
    "tender": [
        "The counter is a sale without a job: a CounterSale with its lines and one or more CounterSaleTenders, inside a "
        "TillSession that is opened, counted and closed. A gift card is a liability the workshop sells: GiftCard is the "
        "card and GiftCardEntry every movement on it, typed as an issue, a redemption, a refund or an adjustment, with a "
        "database constraint that refuses a kind whose sign disagrees with it. The card's status is never stored; it is "
        "read from its entries and dates.",
        "Money held ahead of a bill is a Deposit of kind deposit or credit, and DepositMovement is every amount that "
        "came in, went against a bill, or went back. What is held is the sum, never a stored figure.",
    ],
    "ledger": [
        "The ledger is double entry. Account is the chart, filled to a standard set for every organisation; "
        "JournalEntry is one balanced posting with a source type and id pointing back at the invoice, payment, expense "
        "or bill that caused it, and JournalLine is each debit or credit against an account. A posting that does not "
        "balance is refused before it is written. A reversal is a new entry that points at the one it reverses.",
        "FiscalPeriod is what a bookkeeper closes; a closed period refuses postings. Expense is money spent outside a "
        "supplier bill, claimed and approved; BankAccount, BankStatement and BankStatementLine are the imported bank feed "
        "and its matching against ledger entries, line by line, with the reason recorded.",
    ],
    "stock": [
        "Product is a part or a stocked item with a category; Warehouse is where stock sits (a site or a van); "
        "StockLevel is the quantity on hand per product per warehouse, and StockMovement is every change to it with a "
        "reason. StockReservation holds parts for a job from the moment they are on it, and is consumed when the job "
        "completes, so two jobs cannot both count the same filter.",
        "StockTransfer moves stock between warehouses, dispatched at one and received at the other. Warranty and "
        "WarrantyClaim are the cover on a part or a job and the claim against it when it fails.",
    ],
    "purchasing": [
        "A PurchaseRequest is raised from a job's part line by whoever needs it and approved by whoever may approve "
        "spending; approved requests become a PurchaseOrder to a Supplier, line by line. GoodsReceipt is what actually "
        "arrived, which raises stock levels and writes movements; SupplierBill is what the supplier then asks for, "
        "posted as a payable; SupplierPayment settles bills, allocated across them in SupplierPaymentAllocation.",
        "Price history lives on the order lines, so the product page can say what this part cost last time and from whom.",
    ],
    "people": [
        "Employee is a person on the payroll, which is not the same as a User who can sign in: a technician may be both, "
        "an owner may be neither. PayRate is what they are paid and from when; AttendanceRecord, Timesheet and Absence "
        "are the hours; Skill, EmployeeSkill, JobTypeSkill and Certification are what they can do and what a job needs, "
        "which the planner uses to refuse a job to a bay that cannot take it. WorkingPattern is the week they work.",
        "PayrollRun is one pay period, computed and approved by a role that is deliberately separate from operational "
        "reporting, because the manager who runs the productivity board must not see what the person beside them is "
        "paid. CommissionScheme, CommissionBand and CommissionAward are earnings on work sold.",
    ],
    "communication": [
        "Everything sent to a customer is a MessageLog row before it is a message: the channel, the body, the reason, and "
        "then a DeliveryEvent for every state the provider reports back. InboundMessage is the reply, matched to the job "
        "it is about; SuppressedAddress is where the product stops sending because something bounced or somebody asked.",
        "Retention is ReviewRequest (a strong rating invited to a public review, a poor one into a private conversation), "
        "LoyaltyEntry (points earned and spent, mirrored in the ledger as a liability) and Referral (a code, the person it "
        "brought, and the reward settled when their first bill is paid). AutomationRule, AutomationRun and "
        "AutomationAction are the workshop's own rules and every time they fired; Campaign and CampaignRecipient are "
        "one-off sends to a segment. AssistantExchange is every question asked of the assistant and the answer given, "
        "with the data it was allowed to see.",
    ],
}

SHOT_CAPTIONS: list[tuple[str, str, str]] = [
    ("home", "The public front door.", "Public"),
    ("signup", "Sign-up, with the country that sets the currency, locale, timezone and first tax rate.", "Public"),
    ("login", "Sign in.", "Public"),
    ("dashboard", "The owner's dashboard: today's work, approvals waiting, money in and out.", "Workshop"),
    ("dashboard-dark", "The same, in the dark theme.", "Workshop"),
    ("diary", "The diary: bookings against bays with the capacity check on the drop.", "Workshop"),
    ("job", "The job: the central record, with lines, time, photographs and the timeline.", "Workshop"),
    ("inspection", "Inspection capture, built for a tablet with gloves on.", "Shop floor"),
    ("estimate-approval", "The estimate as the customer approves it, on their phone.", "Customer"),
    ("invoice", "An invoice, with what is owed and every way to take it.", "Workshop"),
    ("receipt", "A receipt for one payment, as it prints.", "Workshop"),
    ("counter", "The counter: one sale on one screen.", "Workshop"),
    ("gift-cards", "Gift cards sold and outstanding.", "Workshop"),
    ("card-payments", "Every card payment tried through the provider, with the provider's record against the books.", "Workshop"),
    ("customer", "A customer, with their vehicles, bills, money held and points.", "Workshop"),
    ("vehicle", "A vehicle as the workshop sees it.", "Workshop"),
    ("passport", "The vehicle passport: the hash-chained history, verified on the page.", "Workshop"),
    ("passports", "Every passport the workshop has touched, with its chain verdict.", "Workshop"),
    ("floor", "The shop floor board.", "Shop floor"),
    ("shop", "The technician's own surface.", "Shop floor"),
    ("people", "People, skills and the week.", "Workshop"),
    ("plan", "The planner: capacity against demand.", "Workshop"),
    ("business", "Business reporting.", "Workshop"),
    ("reputation", "Reviews and reputation.", "Workshop"),
    ("ask", "Ask: a question in words, answered from the workshop's own data.", "Workshop"),
    ("settings-organisation", "Organisation settings, where the region is changed.", "Workshop"),
    ("terms", "The terms: the trial, the charge that follows it, cancelling, the data.", "Public"),
    ("pricing", "The pricing page in the currency of the country picked.", "Public"),
    ("subscription", "The billing page: the card, the next charge as a date and an amount, the invoices.", "Workshop"),
    ("data", "Data in and out: import a CSV, export every dataset.", "Workshop"),
    ("portal", "The customer's home on a phone: what needs them first, then their cars.", "Customer"),
    ("portal-desktop", "The same home on a desktop, two columns under a persistent navigation.", "Customer"),
    ("portal-rewards", "Rewards: points per workshop and the referral code.", "Customer"),
    ("portal-workshops", "The directory of every workshop on Carnet, with Book here on each.", "Customer"),
    ("portal-settings", "The customer's account: details, password, consent per workshop, sign out.", "Customer"),
    ("portal-invoices", "The customer's invoices, with Pay now on what is owed.", "Customer"),
    ("portal-passport", "A passport as its owner sees it.", "Customer"),
    ("portal-passports", "Every car the customer owns, listed.", "Customer"),
    ("not-found", "Not found, in the product's own voice.", "Public"),
    ("mark", "The mark.", "Brand"),
]

VERIFY_TEXT = {
    "verify-authz": "every workshop page as every role, every server action read for its permission, another workshop's records by id, the machine routes with nothing in hand, a read-only workshop, the plan at the point of use",
    "verify-controls": "every control on every page does something and every link goes somewhere",
    "verify-access": "labels, names, alternatives, heading order, landmarks and skip links on every page",
    "verify-empty": "every page works with nothing in it, and says what to do next",
    "verify-ledger": "postings balance, the trial balance and the statements agree",
    "verify-money": "the invoice, payment, credit note and refund cycle against a real database",
    "verify-integrity": "the passport chain: append, verify, tamper, detect",
    "verify-network": "what another workshop may read, and what it may not",
    "verify-offline": "the technician's queue applied through /api/sync, idempotently",
    "verify-region": "a workshop signed up in Canada and moved to France, every figure in the right currency",
    "verify-deposits": "money held: taken, applied, returned, kept as credit, and every posting",
    "verify-giftcards": "gift cards sold, spent at the counter, refunded onto the card, lapsed and cancelled",
    "verify-loyalty": "points earned on payment, spent on a bill, mirrored in the ledger",
    "verify-referrals": "a code redeemed at sign-up and the reward settled on the first paid bill",
    "verify-backup": "a backup taken and restored, and the restored database answering",
    "verify-headers": "the security headers and the content security policy on every surface",
    "verify-import": "five files in three delimiters uploaded, mapped from their headers, previewed with every problem named by row, imported through the pages' own rules, and every dataset exported and read back",
    "verify-billing": "a workshop opens, accepts the terms, adds a card, is reminded once, is charged through the signed webhook, fails and recovers, cancels on a plan and on a trial",
    "verify-payments": "a card payment from the portal: declined with a reason, paid once however often the provider says so, refunded through the provider, reconciled, and the webhook over HTTP",
    "verify-arrival": "a car checked in with every refusal tried, the keys moved, a quote approved with a signature, the car refused unfinished and unpaid, handed back signed for, and a fleet car let go owing with the reason written down",
    "verify-workflow": "a job walked from a booking to a paid bill the way a front desk would, every hand-off counted in presses against the table in core, the three screens read for the one button each needs, and stock put back as it was",
}


# ----------------------------------------------------------------- build
def build() -> None:
    schema = json.loads((ASSETS / "schema.json").read_text(encoding="utf-8"))
    doc = Document()
    style_document(doc)

    # Cover
    para(doc, "CARNET", size=11, colour=BRASS, bold=True, font=DISPLAY, space_after=60)
    para(doc, "System Architecture Document", size=30, colour=PETROL_DEEP, bold=True, font=DISPLAY, space_after=10)
    para(doc, "How the system is put together, how it processes data, and what the database looks like — "
              "with the interface photographed from the running application.", size=13, colour=FAINT, space_after=40)
    para(doc, f"Built on {BUILT_ON}, after {AFTER_PHASE}.", size=9.5, colour=FAINT)
    para(doc, f"{schema['counts']['models']} tables, {schema['counts']['enums']} enumerations, "
              f"{schema['counts']['fields']} columns; {len(list(MIGRATIONS.glob('*/migration.sql')))} migrations.",
         size=9.5, colour=FAINT)
    para(doc, "Regenerated with every change: npm run docs:all. Diagrams are read from the Prisma schema, "
              "screenshots from the running app, counts from the code. Nothing in this document is typed twice.",
         size=9.5, colour=FAINT, italic=True)
    page_break(doc)

    # Contents
    doc.add_heading("Contents", level=1)
    for line in [
        "1. Purpose, and how to read this document",
        "2. The system in one picture",
        "3. Runtime architecture",
        "4. How a request is processed",
        "5. How data is written",
        "6. The workflows, as data flows",
        "7. Database architecture, domain by domain",
        "8. Security in the data path",
        "9. The user interface: wireframes and photographs",
        "10. What verifies it",
        "Appendix A. Every table, by domain",
        "Appendix B. Migrations",
        "Appendix C. Rebuilding this document",
    ]:
        para(doc, line, space_after=2)
    page_break(doc)

    # 1
    doc.add_heading("1. Purpose, and how to read this document", level=1)
    para(doc, "This is the one document that explains how Carnet works as a system: what runs where, what happens "
              "to a request from the moment a browser sends it to the moment a row is written, how the twelve domains "
              "of the database fit together, and what the person at the screen sees while all of that happens. It is "
              "written for an engineer joining the project, a reviewer deciding whether to trust it, and a buyer who "
              "wants to know what they are buying.")
    para(doc, "It is generated, not maintained by hand. The entity-relationship diagrams are drawn from "
              "packages/db/prisma/schema.prisma by infra/scripts/architecture-diagrams.ts; a model added to the schema "
              "appears in the next build, and one missing from the domain table fails the build rather than going "
              "quietly missing. The screenshots are taken from the running application by infra/scripts/design-shots.ts. "
              "The prose lives in infra/scripts/build_architecture_doc.py beside the code it describes and is revised "
              "with each phase.")
    para(doc, "Where this document says how something works, the source of truth is named so it can be checked: a file, "
              "a function, or a design document in docs/02-design. Where those documents go deeper — the security "
              "model, the accounting model, the state machines — this one summarises and points.")

    # 2
    doc.add_heading("2. The system in one picture", level=1)
    diagram(doc, "flow-context", "Who uses Carnet, and what it depends on.", max_h=11)
    para(doc, "Five kinds of people reach the system, and each has a surface of their own rather than the same "
              "application with items greyed out (ADR-0009). Workshop staff use the workshop at /dashboard and its "
              "siblings. A technician uses /shop on a tablet, and it keeps working out of signal. A vehicle owner uses "
              "/portal on a phone. A buyer opens a passport link at /p/<token> with no account at all. Another system "
              "reads /api/v1 with an organisation key.")
    para(doc, "Every external dependency is optional to the core workflow. If email, SMS or VIN decoding is down, a "
              "workshop can still book, inspect, estimate, repair, invoice and take payment. The only thing the system "
              "cannot run without is PostgreSQL.")

    # 3
    doc.add_heading("3. Runtime architecture", level=1)
    diagram(doc, "flow-containers", "The containers: one web application, three packages, a worker for development, one database.", max_h=13)
    para(doc, "Carnet is a monorepo with one deployable web application and three packages it is built from "
              "(ADR-0003). The rules live in packages/core, which imports nothing from the rest of the repository: no "
              "ORM, no React, no framework. That is what makes the domain testable without a database — 1,322 tests "
              "run in under a second — and portable if the framework changes. packages/db is the Prisma client with "
              "the tenant scope built in; packages/ui is the design system, plain React with no data access.")
    doc.add_heading("Layering", level=2)
    table(doc, ["Layer", "May import", "Must not import"], [
        ["packages/core", "nothing from this repository", "Prisma, React, Next"],
        ["packages/db", "core", "React, Next"],
        ["packages/ui", "core, React", "Prisma, Next"],
        ["apps/web", "every package", "—"],
        ["apps/worker", "core, db", "React, ui"],
    ], widths=[4.0, 6.0, 6.0])
    doc.add_heading("Technology", level=2)
    table(doc, ["Concern", "Choice", "Why"], [
        ["Web framework", "Next.js, app router, server components", "The workshop UI ships almost no JavaScript, which is the difference between a list and a spinner on a tablet over patchy Wi-Fi."],
        ["Writes", "Server actions, one per form", "The form posts to the function that owns the rule; no API layer to keep in step with the page."],
        ["Database", "PostgreSQL through Prisma", "Row locks, transactions, check constraints and JSON where the shape varies."],
        ["Rules", "TypeScript in packages/core, tested with Vitest", "Pure functions over plain values; no test touches the database."],
        ["Sessions", "First-party, a row per session", "Revocable on the server; a stolen database yields no usable cookie (ADR-0006)."],
        ["Email", "Resend, with inbound webhooks", "Delivery events and replies come back to the job they are about."],
        ["Files", "Object storage behind signed URLs", "Photographs and documents never pass through the application server twice."],
        ["Scheduled work", "HTTP endpoints called by a scheduler", "The application is hosted where a loop cannot be held open; the worker runs the same functions in development."],
        ["Verification", "tsx scripts against a built server and a real database", "Fifty scripts, one per capability, each cleaning up after itself."],
    ], widths=[3.0, 5.0, 8.6])
    doc.add_heading("Repository layout", level=2)
    mono(doc, "\n".join([
        "apps/web/src/app/(workshop)/*      the workshop pages and their server actions",
        "apps/web/src/app/shop, portal, p   the shop floor, the customer's side, the public passport",
        "apps/web/src/app/api/*             sync, cron, v1, uploads, inbound webhooks, reports",
        "apps/web/src/lib/*                 one module per capability: money.ts, jobs.ts, passport.ts, billing.ts …",
        "apps/web/src/lib/guard.ts          requireSession(permission): the one door every page goes through",
        "packages/core/src/*                the rules, by domain: accounting, billing, passport, region, giftcards …",
        "packages/db/prisma/schema.prisma   the database, and migrations/ beside it, hand-written SQL",
        "packages/db/src/tenant.ts          forOrg(): the tenant scope",
        "packages/ui/src/*                  the design system",
        "infra/scripts/verify-*.ts          the verification scripts",
        "docs/                              product, requirements, design, ADRs, engineering, project",
    ]))

    # 4
    doc.add_heading("4. How a request is processed", level=1)
    para(doc, "Every page and every server action in the workshop goes through one function, requireSession, and "
              "that function is where the system decides who is asking, what they may do, where their workshop stands, "
              "and which rows they may see. The diagram is the whole path; the text after it is each step in detail.")
    diagram(doc, "flow-request", "A request, from cookie to redirect. apps/web/src/lib/guard.ts and auth.ts.", max_h=22)
    doc.add_heading("Step by step", level=2)
    bullets(doc, [
        "The cookie. The browser sends carnet_session, a random 256-bit token. It is never stored: its SHA-256 digest is "
        "the key of the Session row. currentUser() hashes it, finds the session, checks it has not expired, and loads the "
        "user with their membership, the organisation, the organisation's region and its four subscription dates. This "
        "runs once per request and is cached for the rest of it with React's cache().",
        "The subscription. The organisation's state is read from its dates — trial, grace, on a plan, read-only, "
        "cancelled — and carried on the session. If the workshop is read-only and the permission asked for is a writing "
        "one (anything ending in write, complete, send, approve, request, transfer, share or manage), the request is sent "
        "to /settings/subscription with the reason before anything else is looked at.",
        "The permission. can(role, permission) is deny-by-default: a permission not explicitly granted to a role is "
        "refused. A refusal does not throw, because by then the page has started streaming and a thrown error would "
        "arrive as a 200 with a stack trace behind it; it redirects to the role's own surface with ?denied= and the "
        "permission, and the refusal is audited. A technician never reaches a workshop page at all: the workshop layout "
        "turns them back to /shop first.",
        "The scope. forOrg(orgId) returns a Prisma client where every tenant-scoped model has the organisation "
        "predicate added to reads and stamped on creates; findUnique is rewritten to findFirst with the predicate, so an "
        "id lookup cannot escape the tenant. The page receives this client as db and there is no convenient way to get "
        "an unscoped one, so the safe path is also the shortest. A row belonging to another organisation is not found, "
        "never refused, so its existence is not disclosed.",
        "The page. A server component queries through db, derives anything computed through packages/core (totals, "
        "statuses, balances are never stored, ADR-0011), formats money and dates from the organisation's region, and "
        "renders on the server. The browser receives HTML.",
        "The action. A form posts to a server action in the same file. The action calls requireSession with the "
        "permission the write needs, calls one function in apps/web/src/lib that owns the rule, and redirects back to the "
        "page with ?done=, ?saved= or ?problem= so the outcome is a sentence on the screen and the URL is bookmarkable. "
        "A rule that refuses throws a typed error with a sentence in it; the action catches that type and shows the "
        "sentence, and lets anything else surface.",
        "The plan. Inside the library function, not on the page, assertPlanAllows and assertWithinLimit refuse a "
        "feature the plan lacks or one more of a thing the plan counts, naming the plan that would allow it. The page's "
        "own wall is a courtesy; this is the lock.",
    ])
    doc.add_heading("The other surfaces", level=2)
    para(doc, "The customer portal goes through requireUser rather than requireSession: a portal user has no "
              "organisation, and what they may see is decided by ownership (VehicleOwnership) rather than membership. "
              "A customer may book at any workshop on Carnet: the directory lists every active organisation with a "
              "site, and a first booking creates their Customer record there, so the workshop sees a person in its diary. "
              "The public passport at /p/<token> has no session at all; the token is the PassportShare row and its scope "
              "and expiry decide what is rendered, and every opening is recorded. /api/v1 authenticates an organisation "
              "key (hashed at rest, shown once) and reuses the same library functions the pages use, so the API and the "
              "page cannot disagree about what a workshop may see. /api/cron requires a bearer secret compared in "
              "constant time, and refuses if no secret is configured, so a misconfigured deployment fails closed.")

    # 5
    doc.add_heading("5. How data is written", level=1)
    diagram(doc, "flow-write", "A consequential write: strongly consistent inside the transaction, eventually consistent after it.", max_h=13)
    para(doc, "Completing a job is the most consequential write in the system, and its shape is the shape of every "
              "write that has side effects. Inside one transaction the status changes, the JobStatusEvent is appended, "
              "and an OutboxEvent is written saying what happened. The commit is the only thing that must succeed, "
              "because a workshop losing a completed job is unacceptable. Everything downstream — passport events, "
              "stock reservations consumed, ledger entries posted, the customer told — is done by dispatchOutbox, "
              "at-least-once, with an idempotency key on every effect so a retry cannot post twice (ADR-0007).")
    doc.add_heading("The conventions every write follows", level=2)
    bullets(doc, [
        "Money is Decimal(12,2) in the database and integer minor units in the rules. fromDecimal converts at the "
        "boundary, with the currency, and sum() refuses to combine two currencies. No arithmetic on money happens in "
        "floating point (ADR-0008).",
        "Totals, balances and statuses are derived, never stored, wherever a stored figure could drift from the rows "
        "beneath it: an invoice's total from its lines, a gift card's balance from its entries, money held from its "
        "movements, a subscription's state from its dates (ADR-0011).",
        "Every ledger posting is a balanced JournalEntry with lines, refused before it is written if the debits and "
        "credits differ, and refused if the period is closed. Reversals are new entries that point at what they reverse. "
        "The accounting model in docs/02-design/07-accounting-model.md lists every posting.",
        "Append-only tables are never updated or deleted: VehicleEvent, JobStatusEvent, AuditLog, GiftCardEntry, "
        "DepositMovement, DeliveryEvent, ConsentChange. A correction is a new row that supersedes an old one.",
        "Gapless numbers (invoices, credit notes, purchase orders) come from DocumentSequence under a row lock, one "
        "sequence per organisation and kind.",
        "Concurrency on the passport is a row lock on the vehicle, because two workshops completing jobs on the same "
        "car in the same second must not both claim the same sequence number.",
        "Every consequential action writes an AuditLog row with who, what, before and after. The audit write never "
        "throws: a failure to audit must not roll back the action it describes.",
        "Every message to a customer is a MessageLog row first, so what was said and when is a query, and the send "
        "itself is a side effect that can fail and be retried without saying it twice.",
    ])

    # 6
    doc.add_heading("6. The workflows, as data flows", level=1)
    para(doc, "This section follows the data through the workshop's work: what is created, what it points at, and "
              "what it causes. Each subsection names the tables and the module that owns the rule.")

    doc.add_heading("6.1 From a booking to a paid bill", level=2)
    diagram(doc, "flow-workflow", "The core workflow. Booking → Job → Inspection → Estimate → Approval → Invoice → Payment → Ledger and Passport.", max_h=8)
    para(doc, "A Booking takes a diary slot against a bay after the capacity check in packages/core/src/scheduling "
              "agrees there is room. Converting it makes a Job with the vehicle, the customer and the job types, and a "
              "StockReservation for any part already on it. An Inspection runs the workshop's template against the job, "
              "each item rated and photographed. An Estimate is built from the results with the organisation's "
              "currency, sent to the customer as a link, and approved on their phone, line by line, into "
              "EstimateApproval. The work is done against the clock (TimeEntry) and signed off (QualityCheck). "
              "Completing the job goes through the outbox: passport events are appended for every part fitted and every "
              "service done, reservations are consumed into StockMovements, and the customer is told.")
    para(doc, "An Invoice is raised from the job's lines with the tax rate in force and a gapless number, posted to "
              "the ledger as revenue, tax and a receivable. A Payment settles it by one of the methods — cash and card to "
              "the bank; a card through the provider posts the same as a card at the counter; loyalty points, a gift "
              "card or money held go against their own liabilities — and each posting "
              "releases exactly what it should. recordPayment in apps/web/src/lib/money.ts also awards loyalty points, "
              "settles a referral if this was the referred customer's first paid bill, and sends a receipt. A CreditNote "
              "reverses; a Refund sends money back, or keeps it as credit on the customer's account.")

    doc.add_heading("6.2 Card payment through a provider", level=2)
    diagram(doc, "flow-card", "Pay now, from the portal: the attempt, the provider's checkout, one settlement, and reconciliation.", max_h=7)
    para(doc, "A payment provider sits behind a seam of four calls in packages/core (open a checkout, ask what became "
              "of it, refund, read a webhook), with a stub that takes pretend money at a checkout page of ours and "
              "Stripe written against its HTTP API behind the same seam; PAYMENT_PROVIDER chooses. \"Pay now\" on an "
              "outstanding invoice in the portal opens a PaymentAttempt and sends the customer to the checkout. The "
              "provider's answer arrives by webhook or when the customer returns and the page asks; settleAttempt "
              "claims the move from pending with a conditional update before recording anything, so two answers "
              "record one payment, posted exactly like a card at the counter. A declined card is a row with the "
              "provider's reason. A refund by card goes back through the provider, asked before the refund is recorded. "
              "Reconciliation asks the provider about every attempt of the last ninety days and names each difference "
              "between its record and the books without changing anything; the provider is a witness, never the record.")

    doc.add_heading("6.3 The counter, gift cards and money held", level=2)
    para(doc, "A counter sale is the same money path without a job: a CounterSale with lines and tenders in one "
              "transaction, stock moved as it is sold, the ledger posted per tender, inside a TillSession that is "
              "counted at close. A gift card is sold at the counter or on its own page, and the sale posts a liability, "
              "not revenue; spending it releases the liability and never touches cash. A refund by gift card goes back "
              "onto the card that paid. Money held — a deposit before the work, or a refund kept as credit — is the same "
              "idea for a customer rather than a card: a liability until an invoice takes it or it goes back, oldest "
              "first, refused rather than clamped. packages/core/src/giftcards and deposits hold the rules; "
              "apps/web/src/lib/giftcards.ts, deposits.ts and counter.ts do the work.")

    doc.add_heading("6.4 Purchasing and stock", level=2)
    diagram(doc, "flow-purchasing", "Request → order → receipt → bill → payment, with stock rising at the receipt.", max_h=6)
    para(doc, "A part a job needs and the workshop does not have becomes a PurchaseRequest, approved by a role that "
              "may approve spending, gathered onto a PurchaseOrder to a supplier. The GoodsReceipt is the moment stock "
              "exists: StockLevel rises, a StockMovement records it with the order as its reason, and the price paid is "
              "kept on the line for the product's history. The SupplierBill posts a payable; SupplierPayment settles "
              "bills and is allocated across them. Transfers between sites are dispatched at one warehouse and received "
              "at another, and the stock is in transit between.")

    doc.add_heading("6.5 The passport", level=2)
    diagram(doc, "flow-passport", "An event on a vehicle: chained, verified, shared, transferred, read by the network.", max_h=12)
    para(doc, "The passport is the product's reason to exist, and the mechanism is small. Every event on a vehicle is "
              "a VehicleEvent whose hash covers its own content and the hash of the event before it "
              "(packages/core/src/passport/chain.ts). Altering any event changes its hash and breaks every hash after "
              "it; verifyVehicleChain recomputes the chain and names the first divergence, and the passport page shows "
              "the verdict. It is deliberately not a blockchain (ADR-0005): the question is whether the operator of this "
              "database can silently rewrite the past, and a hash chain with periodic anchoring answers it.")
    para(doc, "A share is a PassportShare: a link with a scope (what kinds of events) and an expiry, every opening "
              "recorded. A transfer of ownership is an invitation to the new owner; the history stays, the previous "
              "owner's private notes do not. Another workshop reads what it is entitled to through VehicleAccessGrant, "
              "earned by working on the car or granted by the owner, and the same entitlement governs /api/v1.")

    doc.add_heading("6.6 The technician's device, out of signal", level=2)
    diagram(doc, "flow-sync", "The queue empties through /api/sync when the device is back in signal.", max_h=5)
    para(doc, "The shop floor keeps working without a connection. Actions taken while away — a task ticked, time "
              "started, a note — are queued on the device by a service worker and sent to /api/sync when signal "
              "returns. The route is authenticated like every other (a queue that accepted work on the strength of an id "
              "in the body would be a write API with no door on it), applies each action once by its client-side id "
              "(SyncedAction), and answers with everything the device needs to carry on: its jobs, tasks and parts.")

    doc.add_heading("6.7 Scheduled work", level=2)
    diagram(doc, "flow-scheduled", "Two endpoints, called on a schedule, run the same functions the worker runs on a loop.", max_h=7)
    para(doc, "There is no queue server. A scheduler calls /api/cron/dispatch every minute to drain the outbox and "
              "/api/cron/hourly to queue service reminders, close abandoned time entries, run due automation actions and "
              "audit passport chains. In development apps/worker calls the same functions on a five-second loop. There "
              "are not two implementations of any job.")

    doc.add_heading("6.8 Communication, in and out", level=2)
    para(doc, "Every outbound message is a MessageLog row with a channel and a reason, written before it is sent, then "
              "sent through the provider, whose webhooks come back as DeliveryEvents. A reply arrives at /api/inbound as "
              "an InboundMessage and is matched to the job it is about, so the conversation is on the job. A bounce or "
              "an opt-out writes SuppressedAddress and the product stops sending to it. Automation rules and campaigns "
              "produce messages the same way; nothing bypasses the log.")

    doc.add_heading("6.9 Region and currency", level=2)
    para(doc, "The organisation carries one country, currency, locale and timezone, chosen at sign-up and editable by "
              "the owner. The session carries the region and every page formats money and dates from it; invoices and "
              "estimates are raised in the organisation's currency; a message to a customer is formatted the same way. "
              "A driver in the portal sees each amount in the currency of the workshop that raised it. Changing the "
              "currency converts nothing already recorded, and the page says so. The plan itself is priced in the "
              "workshop's currency from a price book in core that carries every plan in each currency a country can "
              "set, and the provider is asked to charge that currency and no other.")

    doc.add_heading("6.10 The subscription and the plan", level=2)
    diagram(doc, "flow-subscription", "The subscription's states, read from four dates on the organisation.", max_h=11)
    para(doc, "Sign-up records the terms accepted and opens the provider's subscription checkout: the card goes on "
              "file before the fourteen-day trial starts, and until it does the workshop is read-only. The trial's end "
              "date is the provider's, because it is the day the provider charges the card; three days before, the "
              "hourly job sends the owner one reminder with the amount and the date. A paid invoice event makes the "
              "workshop active, a failed one marks it past due and tells the owner, cancelling on a plan ends at the "
              "period's end and on a trial ends at once. The state is read from the dates and the provider's last word "
              "each time a session is made and is never "
              "stored. A lapsed or cancelled workshop keeps reading and exporting and loses writing; nothing is deleted "
              "in any state; choosing a plan reopens it. The plan's limits (sites, bays, vehicles) and features (the "
              "ledger, payroll, transfers, the API) are enforced inside the functions that make or do the thing, with "
              "the refusal naming the number and the plan.")

    doc.add_heading("6.11 Data in, data out", level=2)
    para(doc, "A CSV from another system is parsed in core (any delimiter, quoted fields, a byte-order mark) and "
              "kept as it came on an ImportBatch: headers, rows, the delimiter, and a mapping suggested from the headers "
              "and edited by a person. Validation runs every row through the dataset's rules before anything is written "
              "and names each problem by spreadsheet row and field. The import writes what passed one row at a time "
              "through the same library functions the pages use (createCustomer, registerVehicle or adoptVehicle, "
              "createProduct with an opening stock adjustment, createSupplier, postEntry for a balanced opening entry), "
              "so tenant scoping, duplicate refusal and plan limits apply as they would to typed rows, and a report of "
              "what did not go in is kept on the batch. Exports read each dataset whole through the scoped client and "
              "write CSV in the same shapes, the passports as JSON in the network's form, or everything as one file; "
              "they need only the read permission a read-only or cancelled workshop keeps.")

    doc.add_heading("6.12 The workflow, as one thing", level=2)
    para(doc, "Phase 34 joined the seams between the objects above so that the five workflows run without "
              "retyping. startJobFromBooking creates the job through createJob with everything the booking knew "
              "and marks the booking CONVERTED with the job's id in one transaction. completeInspection raises the "
              "estimate and calls retotalEstimate, which prices every line and the totals from the same "
              "calculateTotals the invoice uses. approveInPerson marks a draft sent and records the decision through "
              "the same recordDecision the customer's page and the share link call, with the staff member named on "
              "the approval row. recordDecision copies the approved lines onto the job and, for each part line that "
              "names a product, calls reserveForJobLine against the default warehouse after the transaction commits, "
              "so the reservation or the purchase request for the shortfall exists the moment the approval does. "
              "quoteExtraWork opens or extends a supplementary estimate as the job's next version; sendEstimate and "
              "recordDecision leave the job's status alone for a supplementary estimate, so the work already approved "
              "carries on and the approved extra lines are appended after the lines already there. The press count "
              "per hand-off is data in core (WORKFLOWS), held by a unit test and by the walkthrough script, and the "
              "table below is generated from it.")
    workflows_table(doc)

    doc.add_heading("6.13 Arrival and departure", level=2)
    para(doc, "checkIn writes, in one transaction, the VehicleCheck (IN) with the mileage, fuel, marks, belongings "
              "and hook, the Signature it points at, a KeyEvent (HUNG) if a hook was given, the OdometerReading "
              "on the vehicle and the job's receivedAt; the rules (assertCheckInAllowed, assertSigned) run in core "
              "first, so a mileage that goes backwards or an empty pad is refused before anything is written. "
              "Photographs go to the store on a ticket scoped to the check and are recorded as JobAttachments with "
              "checkId once the store confirms them. checkOut reads what is outstanding across the job's live "
              "invoices, refuses a car that is not completed or that owes money with no reason given, then writes "
              "the VehicleCheck (OUT), its Signature, a KeyEvent (HANDED_TO_CUSTOMER), the final reading and "
              "Job.collectedAt. moveKeys appends a KeyEvent and the present whereabouts is keyState over the "
              "events, read by the job page, the technician's card and the portal. recordSignature writes a "
              "Signature for an approval given across the counter, which EstimateApproval.signatureId points at.")

    doc.add_heading("6.14 Answers and the assistant", level=2)
    para(doc, "Ask takes a question in words and answers from the workshop's own data through the same scoped client "
              "the pages use, so it cannot see more than the person asking. Every exchange is an AssistantExchange row "
              "with the data it was allowed to see. Generated text inside the workflow (a summary for the customer, a "
              "reply drafted) is always a draft a person sends, never a decision (ADR-0013).")

    # 7 — database
    doc.add_heading("7. Database architecture, domain by domain", level=1)
    counts = schema["counts"]
    para(doc, f"One PostgreSQL database, {counts['models']} tables in {len(schema['domains'])} domains, "
              f"{counts['enums']} enumerations, {counts['fields']} columns, defined in packages/db/prisma/schema.prisma "
              f"and applied by {len(list(MIGRATIONS.glob('*/migration.sql')))} hand-written migrations. Every table "
              "that belongs to a workshop carries orgId, and the tenant scope in packages/db/src/tenant.ts adds it to "
              "every query. The tables that are deliberately global — the vehicle and its history, the user, recalls, "
              "part fitments — are the ones the network is built on.")
    doc.add_heading("Conventions", level=2)
    bullets(doc, [
        "Identifiers are cuids, generated by the application, never sequential, never guessable.",
        "Every row has createdAt; rows that change have updatedAt; rows that are retired have isActive or a dated "
        "column (revokedAt, disabledAt, cancelledAt) rather than being deleted.",
        "Foreign keys to the organisation cascade; foreign keys between records restrict, so a customer with invoices "
        "cannot be deleted by accident.",
        "Money is Decimal(12,2) with the currency beside it where it can differ; quantities are Decimal where parts "
        "are sold by measure.",
        "Enumerations are Postgres enums, so an invalid status cannot be written; state is read from facts where a "
        "stored status could drift.",
        "Check constraints guard what the application must never do even if it tried: a gift card entry whose sign "
        "disagrees with its kind, a passport event out of sequence.",
        "Unique constraints carry the business rule: one membership per user per organisation, one vehicle per VIN, "
        "one supplier name per organisation, one number per document sequence.",
    ])
    para(doc, "In the diagrams that follow, PK is the primary key, FK a foreign key, UK a unique column; a column "
              "marked optional is nullable. A box that reads \"in another domain\" is a table drawn in full in another "
              "section. Tables with more than fourteen columns show the first fourteen and the count of the rest.",
          size=9.5, colour=FAINT, italic=True)

    for d in schema["domains"]:
        section = doc.add_section(WD_SECTION.NEW_PAGE)
        landscape(section)
        doc.add_heading(f"7.{schema['domains'].index(d) + 1} {d['title']}", level=2)
        for text in DOMAIN_TEXT.get(d["key"], []):
            para(doc, text, size=10)
        table(doc, ["Table", "Columns", "Tied to a workshop", "Points at"], [
            [m["name"], str(m["fields"]), m["scope"], ", ".join(dict.fromkeys(m["relations"])) or "—"]
            for m in d["models"]
        ], widths=[5.0, 2.0, 4.0, 14.0], size=8)
        page_break(doc)
        diagram(doc, f"erd-{d['key']}", f"Entity-relationship diagram: {d['title'].lower()}. Read from schema.prisma.", max_w=LANDSCAPE_W, max_h=16.2)

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    portrait(section)

    # 8
    doc.add_heading("8. Security in the data path", level=1)
    para(doc, "docs/02-design/05-security-model.md is the full model; this is where it touches the data path.")
    bullets(doc, [
        "Authentication is a row per session, found by the digest of the cookie, expiring, revocable. Passwords are "
        "hashed with a modern KDF; LoginAttempt rate-limits by address and by account; AuthToken carries the one-time "
        "links for sign-up confirmation, invitations and password resets.",
        "Authorisation is three layers, all of which must pass: the permission (deny by default), the scope (the row "
        "must be in the caller's tenant, or entitled by ownership or grant), and the plan (the feature or the room for "
        "one more). A cross-tenant request by id is not found.",
        "Tenant isolation is in the data layer (ADR-0010): the scoped client is the only one a page receives. Its known "
        "gap is stated: update and delete by id are not narrowed, so the convention is to load through the scoped "
        "client first, and row-level security in Postgres is the intended successor.",
        "Machine routes fail closed: /api/cron refuses when no secret is configured; /api/v1 keys are hashed at rest; "
        "/api/sync needs the same cookie as a page.",
        "Public links carry their own scope and expiry and record every opening; the passport exception (a workshop "
        "reading a car it did not own) is an explicit grant, never an absence of a check.",
        "The response carries a content security policy with a per-request nonce and the usual headers, checked by "
        "verify-headers.ts on every surface.",
        "Authorisation is tried, not trusted: verify-authz.ts opens every workshop page as every role and reads every "
        "server action for its permission before each phase ships.",
    ])

    # 9 — UI
    doc.add_heading("9. The user interface: wireframes and photographs", level=1)
    para(doc, "The wireframes are the intent, drawn before the screens were built and kept in "
              "docs/02-design/09-wireframes.md. The photographs are the built screens, taken from the running "
              "application by infra/scripts/design-shots.ts on the day this document was built. The design book "
              "(docs/Carnet-Design-Book.pdf) explains each design; this section shows what the system looks like "
              "while it does what the sections above describe.")
    doc.add_heading("9.1 Wireframes", level=2)
    for title, caption in [
        ("Workshop shell", "The workshop shell: groups named for what a person is doing; items a role cannot reach are absent."),
        ("Job detail", "The job: the central screen."),
        ("Vehicle passport", "The passport, as the workshop sees it."),
        ("Inspection capture", "Inspection capture on a tablet."),
    ]:
        block = wireframe(title)
        if block:
            label(doc, caption)
            mono(doc, block, size=6.8)
    doc.add_heading("9.2 Photographs of the running application", level=2)
    current = None
    for name, caption, surface in SHOT_CAPTIONS:
        if surface != current:
            doc.add_heading(surface, level=3)
            current = surface
        shot(doc, name, caption, width_cm=PORTRAIT_W if name not in ("mark",) else 4.0)

    # 10 — verification
    doc.add_heading("10. What verifies it", level=1)
    para(doc, "Every phase ships behind the same checks: the type checker, the linter, the formatter, the unit tests "
              "in packages/core, a production build, and then a built server crawled by three scripts and exercised by "
              "the phase's own verification script against a real database. The scripts are in infra/scripts and each "
              "removes what it made.")
    scripts = sorted(p.stem for p in SCRIPTS.glob("verify-*.ts"))
    table(doc, ["Script", "What it exercises"], [
        [s, VERIFY_TEXT.get(s, "the capability it is named for, end to end, against a real database")] for s in scripts
    ], widths=[4.5, 12.1], size=8)

    # Appendix A
    page_break(doc)
    doc.add_heading("Appendix A. Every table, by domain", level=1)
    for d in schema["domains"]:
        label(doc, d["title"])
        para(doc, ", ".join(m["name"] for m in d["models"]), size=9, font=MONO)

    # Appendix B
    doc.add_heading("Appendix B. Migrations", level=1)
    para(doc, "Hand-written SQL, applied in order by prisma migrate deploy. The name says what it did.", size=9.5, colour=FAINT)
    table(doc, ["Migration", "Applied as"], [
        [p.parent.name.split("_", 1)[1] if "_" in p.parent.name else p.parent.name,
         f"{p.parent.name[:4]}-{p.parent.name[4:6]}-{p.parent.name[6:8]}"]
        for p in sorted(MIGRATIONS.glob("*/migration.sql"))
    ], widths=[10.0, 6.6], size=8)

    # Appendix C
    doc.add_heading("Appendix C. Rebuilding this document", level=1)
    mono(doc, "\n".join([
        "npm run build --workspace=@carnet/web && npm run start --workspace=@carnet/web   # a built server on :3000",
        "npm run docs:shots         # photographs of the running app  -> docs/assets/design/*.png",
        "npm run docs:diagrams      # ER and flow diagrams from the schema -> docs/assets/architecture/*.png",
        "npm run docs:architecture  # this document                     -> docs/Carnet-System-Architecture.docx / .pdf",
        "npm run docs:all           # the deck, the business case, the design book and this document",
    ]), size=7.5)
    para(doc, "A model added to the schema must be placed in a domain in infra/scripts/architecture-diagrams.ts, or the "
              "diagram build fails and says which. A screen added to the app should be added to design-shots.ts and "
              "captioned here. The prose in build_architecture_doc.py is revised with each phase.", size=9.5, colour=FAINT)

    doc.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}")


def to_pdf() -> None:
    if sys.platform != "win32":
        return
    pdf = OUT.with_suffix(".pdf")
    script = (
        "$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
        f"$d = $w.Documents.Open('{OUT}'); $d.SaveAs2('{pdf}', 17); $d.Close(); $w.Quit(); 'pdf'"
    )
    try:
        out = subprocess.run(["powershell", "-NoProfile", "-Command", script], capture_output=True, text=True, timeout=300)
        if "pdf" in out.stdout:
            print(f"wrote {pdf.relative_to(ROOT)}")
        else:
            print("Word did not produce a PDF; the .docx is complete on its own.")
    except Exception:  # noqa: BLE001
        print("Word is not available; the .docx is complete on its own.")


if __name__ == "__main__":
    build()
    to_pdf()
